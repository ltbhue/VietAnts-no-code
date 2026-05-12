from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(r"c:\Users\DELL\Desktop\Đồ án lớn\Đồ án")
SOURCE_MD = ROOT / "docs" / "final" / "09-bao-cao-hoan-chinh-co-dau.md"
OUTPUT_DOCX = ROOT / "docs" / "final" / "09-bao-cao-hoan-chinh-co-dau.docx"


def add_paragraph_from_markdown(doc: Document, line: str) -> None:
    text = line.rstrip()
    if not text:
        doc.add_paragraph("")
        return

    if text.startswith("# "):
        doc.add_heading(text[2:].strip(), level=1)
        return
    if text.startswith("## "):
        doc.add_heading(text[3:].strip(), level=2)
        return
    if text.startswith("### "):
        doc.add_heading(text[4:].strip(), level=3)
        return
    if text.startswith("#### "):
        doc.add_heading(text[5:].strip(), level=4)
        return
    if text.startswith("- "):
        doc.add_paragraph(text[2:].strip(), style="List Bullet")
        return
    if text[:2].isdigit() and text[1:3] == ". ":
        doc.add_paragraph(text[3:].strip(), style="List Number")
        return
    if text[:3].isdigit() and text[2:4] == ". ":
        doc.add_paragraph(text[4:].strip(), style="List Number")
        return
    if text.strip() == "---":
        doc.add_paragraph("")
        return

    doc.add_paragraph(text)


def main() -> None:
    md_lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)

    for line in md_lines:
        add_paragraph_from_markdown(doc, line)

    doc.save(OUTPUT_DOCX)
    print(f"Exported: {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
